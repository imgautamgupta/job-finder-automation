import sys
import os
sys.path.insert(0, os.path.abspath("."))

from app.database.database import DatabaseManager
from app.services.job_agent import JobAgentService
from app.services.search_intelligence import SearchIntelligenceService
from app.services.excel_exporter import ExcelExporter

db = DatabaseManager("data/jobs.db")
jobs = db.get_all_jobs()
print("=" * 60)
print("PHASE 14 REAL DATABASE VALIDATION REPORT")
print("=" * 60)

print(f"Total jobs in DB: {len(jobs)}")
relevant_jobs = [j for j in jobs if (j.match_score or 0) >= 60]
print(f"Relevant jobs (>=60): {len(relevant_jobs)}")

recommended_jobs = [
    j for j in jobs
    if j.application_priority and (j.application_priority.value if hasattr(j.application_priority, "value") else str(j.application_priority)) in ("URGENT", "HIGH", "NORMAL")
]
print(f"Recommended jobs: {len(recommended_jobs)}")

# Canonical application definition: status in ("APPLIED", "INTERVIEW", "OFFER")
applications = [j for j in jobs if (j.status.value if hasattr(j.status, "value") else str(j.status)) in ("APPLIED", "INTERVIEW", "OFFER")]
print(f"Applications tracked (canonical): {len(applications)}")
for app in applications:
    print(f"  - Job {app.id}: {app.title} @ {app.company} | Status: {app.status} | Applied at: {app.applied_at}")

# Service instantiation
agent_service = JobAgentService(db=db)
report = agent_service.get_agent_report()

print("\n--- Daily Target Progress ---")
print(f"Daily Target: {report.plan.target_daily}")
print(f"Today's applications: {report.plan.target_applied_today}")
print(f"Daily target remaining: {report.plan.target_remaining}")
print(f"Target progress pct: {report.plan.target_progress_pct}%")
print(f"Active streak: {report.plan.streak_days} days")

print("\n--- Follow-ups ---")
follow_ups = [j for j in jobs if j.follow_up_date]
print(f"Scheduled follow-ups in DB: {len(follow_ups)}")
for fu in follow_ups:
    print(f"  - Job {fu.id}: {fu.title} @ {fu.company} | Follow-up date: {fu.follow_up_date} | Status: {fu.status}")
print(f"Follow-ups due in report: {report.summary.follow_ups_due}")
print(f"Overdue actions in report: {report.summary.overdue_actions}")

print("\n--- Interviews ---")
interviews = [j for j in jobs if (j.status.value if hasattr(j.status, "value") else str(j.status)) == "INTERVIEW" or j.interview_date]
print(f"Scheduled interviews: {len(interviews)}")
if not interviews:
    print("0 scheduled interviews")

print("\n--- Resume Preparation States ---")
for j in jobs:
    prep_st = j.preparation_status or "NOT_REVIEWED"
    print(f"  - Job {j.id}: {j.company} | Status: {j.status} | Prep: {prep_st} | Match: {j.match_score} | Readiness: {j.application_readiness_score} | Resume: {j.resume_alignment_score}")

print("\n--- Search Intelligence Status ---")
search_svc = SearchIntelligenceService(db=db)
s_report = search_svc.get_search_intelligence_report()
print(f"Search Intelligence total_jobs: {s_report.total_jobs}")
print(f"Search Intelligence applied_jobs: {s_report.applied_jobs}")
print(f"Search Intelligence confidence: {s_report.data_confidence.value}")
print(f"Search recommendations: {len(s_report.recommendations)}")
for rec in s_report.recommendations[:3]:
    print(f"  - [{rec.category}] {rec.title}: {rec.recommendation}")

print("\n--- Generated Daily Action Plan (Part 3/4) ---")
print(f"Total actions: {len(report.plan.actions)}")
print(f"Urgent actions ({len(report.plan.urgent_actions)}):")
for a in report.plan.urgent_actions:
    print(f"  - [{a.action_type}] {a.title} (Reason: {a.reason})")

print(f"High actions ({len(report.plan.high_actions)}):")
for a in report.plan.high_actions:
    print(f"  - [{a.action_type}] {a.title} (Reason: {a.reason})")

print(f"Medium actions ({len(report.plan.medium_actions)}):")
for a in report.plan.medium_actions:
    print(f"  - [{a.action_type}] {a.title} (Reason: {a.reason})")

print(f"Low actions ({len(report.plan.low_actions)}):")
for a in report.plan.low_actions:
    print(f"  - [{a.action_type}] {a.title} (Reason: {a.reason})")

# Test Excel export with real database jobs
print("\n--- Excel Export Validation ---")
exporter = ExcelExporter(exports_dir="data/exports")
path = exporter.export(jobs)
print(f"Exported workbook: {path}")

import openpyxl
wb = openpyxl.load_workbook(path, data_only=True)
print(f"Sheets in exported workbook ({len(wb.sheetnames)}): {wb.sheetnames}")
assert "Job Agent" in wb.sheetnames
ws_agent = wb["Job Agent"]
print(f"Job Agent rows: {ws_agent.max_row}")

print("\n" + "=" * 60)
print("REAL DATABASE VALIDATION COMPLETE - ALL ASSERTIONS PASSED")
print("=" * 60)
