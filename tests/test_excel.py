"""Unit tests for the Production-Ready Excel Job Report (Phase 5)."""

from datetime import datetime, timezone
from pathlib import Path
import tempfile
import openpyxl
import pytest

from app.models.job import Job
from app.services.excel_exporter import ExcelExporter, REPORT_COLUMNS, get_priority


@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as tmp:
        yield Path(tmp)


@pytest.fixture
def sample_jobs():
    """Create sample scored jobs spanning different categories, scores, and easy-apply flags."""
    j1 = Job(
        source="linkedin",
        title="Senior Python Developer",
        company="BigTech",
        location="Bangalore",
        url="https://linkedin.com/jobs/1",
        easy_apply=False,
        match_score=95,
        match_category="Excellent",
        matched_roles=["Python Developer"],
        matched_skills=["Python", "FastAPI", "PostgreSQL"],
        matched_locations=["Bangalore"],
        experience_match=True,
    )
    j2 = Job(
        source="linkedin",
        title="Python Developer",
        company="Startup Alpha",
        location="Bangalore",
        url="https://linkedin.com/jobs/2",
        easy_apply=True,  # Same score as j3, but easy_apply=True
        match_score=85,
        match_category="Strong",
        matched_roles=["Python Developer"],
        matched_skills=["Python", "FastAPI"],
        matched_locations=["Bangalore"],
        experience_match=True,
    )
    j3 = Job(
        source="naukri",
        title="Backend Engineer",
        company="Startup Beta",
        location="Remote",
        url="https://naukri.com/jobs/3",
        easy_apply=False,  # Same score as j2, but easy_apply=False
        match_score=85,
        match_category="Strong",
        matched_roles=["Backend Developer"],
        matched_skills=["Python", "Docker"],
        matched_locations=["Remote"],
        experience_match=True,
    )
    j4 = Job(
        source="naukri",
        title="Junior Web Developer",
        company="Web Services Ltd",
        location="Hyderabad",
        url="https://naukri.com/jobs/4",
        easy_apply=True,
        match_score=65,
        match_category="Possible",
        matched_roles=["Web Developer"],
        matched_skills=["JavaScript", "React"],
        matched_locations=["Hyderabad"],
        experience_match=False,
        excluded_reasons=["Location not Bangalore"],
    )
    j5 = Job(
        source="linkedin",
        title="Legacy Cobol Programmer",
        company="Bank Corp",
        location="Mumbai",
        url="https://linkedin.com/jobs/5",
        easy_apply=False,
        match_score=40,  # Below 60, should not appear in either sheet
        match_category="Not Relevant",
    )
    return [j1, j2, j3, j4, j5]


def test_workbook_created_successfully(temp_dir, sample_jobs):
    """Test 1: Output file is created and has valid .xlsx extension."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)
    assert file_path.exists()
    assert file_path.suffix == ".xlsx"


def test_both_sheets_exist(temp_dir, sample_jobs):
    """Test 2: Workbook contains 'Top Matches' and 'All Relevant Jobs'."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    sheet_names = wb.sheetnames
    assert "Top Matches" in sheet_names
    assert "All Relevant Jobs" in sheet_names


def test_top_matches_threshold_and_count(temp_dir, sample_jobs):
    """Test 3: Top Matches contains only jobs with score >= 80."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    ws = wb["Top Matches"]
    # Row 1 is header, data rows are 2..
    scores = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert len(scores) == 3  # j1 (95), j2 (85), j3 (85)
    for score in scores:
        assert score >= 80


def test_all_relevant_jobs_threshold_and_count(temp_dir, sample_jobs):
    """Test 4: All Relevant Jobs contains only jobs with score >= 60."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    ws = wb["All Relevant Jobs"]
    scores = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert len(scores) == 4  # j1 (95), j2 (85), j3 (85), j4 (65); j5 (40) excluded
    for score in scores:
        assert score >= 60


def test_sorting_by_score_descending(temp_dir, sample_jobs):
    """Test 5: Jobs are sorted by score descending."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    ws = wb["All Relevant Jobs"]
    scores = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert scores == sorted(scores, reverse=True)


def test_easy_apply_secondary_sorting(temp_dir, sample_jobs):
    """Test 6: When scores are equal, Easy Apply=Yes comes before Easy Apply=No."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    ws = wb["Top Matches"]
    # Rows with score 85: row 3 and row 4
    r2_score = ws.cell(row=3, column=2).value
    r2_easy = ws.cell(row=3, column=8).value
    r3_score = ws.cell(row=4, column=2).value
    r3_easy = ws.cell(row=4, column=8).value

    assert r2_score == 85
    assert r3_score == 85
    assert r2_easy == "Yes"
    assert r3_easy == "No"


def test_required_columns_exist_in_order(temp_dir, sample_jobs):
    """Test 7: All 17 required columns exist in the exact required order."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    for sheet_name in ("Top Matches", "All Relevant Jobs"):
        ws = wb[sheet_name]
        headers = [ws.cell(row=1, column=c).value for c in range(1, len(REPORT_COLUMNS) + 1)]
        assert headers == REPORT_COLUMNS


def test_job_url_clickable_hyperlink(temp_dir, sample_jobs):
    """Test 8: Job URL cell contains active hyperlink."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    ws = wb["Top Matches"]
    url_cell = ws.cell(row=2, column=11)  # Column 11: Job URL
    assert url_cell.value == "https://linkedin.com/jobs/1"
    assert url_cell.hyperlink is not None
    assert url_cell.hyperlink.target == "https://linkedin.com/jobs/1"


def test_priority_generated_correctly():
    """Test 9: Priority function correctly maps scores to HIGH, MEDIUM, NORMAL."""
    assert get_priority(95) == "HIGH"
    assert get_priority(90) == "HIGH"
    assert get_priority(89) == "MEDIUM"
    assert get_priority(80) == "MEDIUM"
    assert get_priority(79) == "NORMAL"
    assert get_priority(60) == "NORMAL"
    assert get_priority(59) == ""


def test_empty_input_produces_valid_workbook(temp_dir):
    """Test 10: Empty job list produces a valid 2-sheet workbook with headers without crashing."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export([])

    assert file_path.exists()
    wb = openpyxl.load_workbook(str(file_path))
    for sheet_name in ("Top Matches", "All Relevant Jobs"):
        ws = wb[sheet_name]
        headers = [ws.cell(row=1, column=c).value for c in range(1, len(REPORT_COLUMNS) + 1)]
        assert headers == REPORT_COLUMNS
        assert ws.max_row == 1  # Headers only


def test_matched_roles_exported(temp_dir, sample_jobs):
    """Test 11: Matched roles column populated."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    ws = wb["Top Matches"]
    roles_val = ws.cell(row=2, column=12).value
    assert "Python Developer" in roles_val


def test_matched_skills_exported(temp_dir, sample_jobs):
    """Test 12: Matched skills column populated."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    ws = wb["Top Matches"]
    skills_val = ws.cell(row=2, column=13).value
    assert "FastAPI" in skills_val
    assert "Python" in skills_val


def test_matched_locations_exported(temp_dir, sample_jobs):
    """Test 13: Matched locations column populated."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    ws = wb["Top Matches"]
    loc_val = ws.cell(row=2, column=14).value
    assert "Bangalore" in loc_val


def test_experience_match_exported(temp_dir, sample_jobs):
    """Test 14: Experience match exported as Yes/No."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    ws_top = wb["Top Matches"]
    assert ws_top.cell(row=2, column=15).value == "Yes"

    ws_all = wb["All Relevant Jobs"]
    # Row 5 corresponds to j4 with experience_match=False
    assert ws_all.cell(row=5, column=15).value == "No"


def test_excluded_reasons_exported(temp_dir, sample_jobs):
    """Test 15: Excluded reasons exported when present, blank when none."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    ws_all = wb["All Relevant Jobs"]
    # Row 5 (j4) has excluded_reasons
    assert "Location not Bangalore" in ws_all.cell(row=5, column=16).value
    # Row 2 (j1) has no exclusions (blank cell)
    assert ws_all.cell(row=2, column=16).value in (None, "")


def test_workbook_opens_on_top_matches(temp_dir, sample_jobs):
    """Test 16: Active sheet index is 0 (Top Matches)."""
    exporter = ExcelExporter(exports_dir=temp_dir)
    file_path = exporter.export(sample_jobs)

    wb = openpyxl.load_workbook(str(file_path))
    assert wb.active.title == "Top Matches"
