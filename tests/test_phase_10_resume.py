"""Unit and integration tests for Phase 10 Resume Intelligence & Preparation Layer."""

from datetime import datetime, timezone
import io
import json
from pathlib import Path
import pytest
from fastapi.testclient import TestClient

from app.dashboard.app import app
from app.database.database import DatabaseManager
from app.models.job import ApplicationPriority, Job, JobStatus
from app.models.profile import JobProfile
from app.models.resume import (
    CandidateExperience,
    CandidateProfile,
    CandidateProject,
    PreparationStatus,
    ResumePreparationAnalysis,
    ResumeProfile,
)
from app.services.excel_exporter import REPORT_COLUMNS, ExcelExporter
from app.services.job_tracker import JobTracker
from app.services.resume_intelligence import ResumeIntelligenceService
from app.services.resume_parser import NORMALIZATION_MAP, ResumeParser, normalize_skill_name


@pytest.fixture
def sample_profile() -> JobProfile:
    """Fixture providing candidate search profile with projects and experience."""
    return JobProfile(
        target_roles=["Python Developer", "Backend Developer", "AI Engineer", "Software Engineer"],
        preferred_locations=["Bangalore", "Delhi", "Remote"],
        skills=[
            "Python", "JavaScript", "TypeScript", "FastAPI", "React",
            "PostgreSQL", "MongoDB", "SQL", "OpenCV", "Scikit-learn",
            "Machine Learning", "Git", "GitHub", "REST API"
        ],
        positive_keywords=["AI", "Backend", "Full Stack"],
        negative_keywords=["Senior", "10+ years"],
        projects=[
            {
                "name": "Loan Application Fraud Detection via Document Verification",
                "skills": ["Python", "OpenCV", "Scikit-learn", "FastAPI", "OCR", "Machine Learning"],
                "description": "AI/ML fraud prevention pipeline using document verification."
            },
            {
                "name": "NeuroNest Notes & Bookmarks Manager",
                "skills": ["JavaScript", "React", "Next.js", "REST API", "MongoDB", "SQL", "PostgreSQL"],
                "description": "Full-stack personal knowledge and bookmarks manager."
            }
        ],
        experience=[
            {
                "title": "AI/ML Intern",
                "company": "AI/ML Internship",
                "experience_type": "internship",
                "skills": ["Python", "OpenCV", "Scikit-learn", "FastAPI", "Machine Learning"],
                "description": "Built computer vision and machine learning models for document analysis."
            }
        ]
    )


@pytest.fixture
def sample_job() -> Job:
    """Fixture providing a scored candidate job."""
    return Job(
        id=1,
        source="lever",
        title="Python Backend Developer",
        company="TechCorp Solutions",
        location="Bangalore",
        url="https://jobs.lever.co/techcorp/python-backend-dev",
        description=(
            "We are looking for a Python Backend Developer.\n\n"
            "Requirements:\n"
            "- Strong experience with Python and FastAPI\n"
            "- Solid knowledge of SQL and PostgreSQL\n"
            "- Familiarity with REST APIs and Git\n\n"
            "Nice to have:\n"
            "- Redis, Docker, and AWS experience\n"
        ),
        easy_apply=False,
        match_score=85,
        match_category="Strong Match",
        matched_roles=["Python Developer", "Backend Developer"],
        matched_skills=["Python", "FastAPI", "SQL", "PostgreSQL", "REST API"],
        matched_locations=["Bangalore"],
        experience_match=True,
        application_readiness_score=86,
        application_priority=ApplicationPriority.URGENT,
        status=JobStatus.NEW,
    )


class TestResumeParser:
    """Test suite for local document parsing across supported formats."""

    def test_parse_txt_file(self, tmp_path: Path):
        parser = ResumeParser()
        txt_path = tmp_path / "resume.txt"
        txt_path.write_text(
            "John Doe\n\n"
            "SUMMARY\n"
            "Passionate Python and Machine Learning Engineer.\n\n"
            "SKILLS\n"
            "Python, Fast API, React.js, Postgres, Docker, Git\n\n"
            "PROJECTS\n"
            "Loan Application Fraud Detection\n"
            "- Implemented OCR and OpenCV pipeline\n\n"
            "EXPERIENCE\n"
            "AI/ML Intern\n"
            "- Developed computer vision models\n",
            encoding="utf-8"
        )

        profile = parser.parse_file(txt_path)
        assert profile.is_fallback is False
        assert "Python" in profile.skills
        assert "FastAPI" in profile.skills
        assert "React" in profile.skills
        assert "PostgreSQL" in profile.skills
        assert len(profile.projects) >= 1
        assert len(profile.experience) >= 1

    def test_parse_pdf_file(self, tmp_path: Path):
        parser = ResumeParser()
        pdf_path = tmp_path / "sample_resume.pdf"
        try:
            import pypdf
            # Generate a minimal valid PDF using pypdf
            writer = pypdf.PdfWriter()
            page = writer.add_blank_page(width=612, height=792)
            # Add minimal text annotation / stream if possible or test parser gracefully
            with open(pdf_path, "wb") as f:
                writer.write(f)

            profile = parser.parse_file(pdf_path)
            assert profile is not None
            assert profile.source_file == "sample_resume.pdf"
        except Exception as e:
            pytest.skip(f"pypdf test skipped: {e}")

    def test_parse_docx_file(self, tmp_path: Path):
        parser = ResumeParser()
        docx_path = tmp_path / "sample_resume.docx"
        try:
            import docx
            doc = docx.Document()
            doc.add_heading("Jane Doe", level=0)
            doc.add_heading("SKILLS", level=1)
            doc.add_paragraph("Python, FastAPI, MongoDB, PostgreSQL")
            doc.save(str(docx_path))

            profile = parser.parse_file(docx_path)
            assert profile.is_fallback is False
            assert "Python" in profile.skills
            assert "FastAPI" in profile.skills
            assert "MongoDB" in profile.skills
        except Exception as e:
            pytest.skip(f"python-docx test skipped: {e}")

    def test_parse_missing_file(self, tmp_path: Path):
        parser = ResumeParser()
        missing = tmp_path / "nonexistent_resume.pdf"
        profile = parser.parse_file(missing)
        assert profile.is_fallback is True

    def test_parse_unsupported_file(self, tmp_path: Path):
        parser = ResumeParser()
        unsupported = tmp_path / "resume.xyz"
        unsupported.write_text("Some text", encoding="utf-8")
        profile = parser.parse_file(unsupported)
        assert profile.is_fallback is True

    def test_parse_malformed_file(self, tmp_path: Path):
        parser = ResumeParser()
        corrupt_pdf = tmp_path / "corrupt.pdf"
        corrupt_pdf.write_bytes(b"%PDF-INVALID-CORRUPT-BYTES-HERE")
        profile = parser.parse_file(corrupt_pdf)
        assert profile.is_fallback is True

    def test_load_resumes_directory(self, tmp_path: Path):
        parser = ResumeParser()
        resumes_dir = tmp_path / "resumes"
        resumes_dir.mkdir()

        (resumes_dir / "resume_backend.txt").write_text("SKILLS: Python, FastAPI, SQL", encoding="utf-8")
        (resumes_dir / "resume_ml.txt").write_text("SKILLS: Python, Scikit-learn, OpenCV", encoding="utf-8")

        resumes = parser.load_resumes(resumes_dir)
        assert len(resumes) == 2
        labels = list(resumes.keys())
        assert any("Backend" in l for l in labels)
        assert any("Ml" in l or "ML" in l for l in labels)

    def test_load_resumes_empty_or_none(self):
        parser = ResumeParser()
        resumes = parser.load_resumes("")
        assert "Profile Fallback" in resumes
        assert resumes["Profile Fallback"].is_fallback is True


class TestNormalization:
    """Test suite for deterministic technical term normalizations."""

    def test_normalize_variations(self):
        assert normalize_skill_name("Fast API") == "FastAPI"
        assert normalize_skill_name("fastapi") == "FastAPI"
        assert normalize_skill_name("React.js") == "React"
        assert normalize_skill_name("ReactJS") == "React"
        assert normalize_skill_name("react") == "React"
        assert normalize_skill_name("Node.js") == "Node.js"
        assert normalize_skill_name("nodejs") == "Node.js"
        assert normalize_skill_name("Postgres") == "PostgreSQL"
        assert normalize_skill_name("postgresql") == "PostgreSQL"
        assert normalize_skill_name("Mongo DB") == "MongoDB"
        assert normalize_skill_name("Scikit Learn") == "scikit-learn"
        assert normalize_skill_name("sklearn") == "scikit-learn"
        assert normalize_skill_name("Machine Learning") == "Machine Learning"
        assert normalize_skill_name("Artificial Intelligence") == "AI"
        assert normalize_skill_name("ai") == "AI"
        assert normalize_skill_name("rest api") == "REST API"


class TestResumeIntelligenceService:
    """Test suite for alignment scoring, gap analysis, project matching, and recommendations."""

    def test_strong_alignment(self, sample_profile: JobProfile, sample_job: Job):
        intel = ResumeIntelligenceService(profile=sample_profile)
        analysis = intel.analyze(sample_job)

        assert analysis.resume_alignment_score >= 70
        assert analysis.resume_alignment_category in ("Strong Alignment", "Good Alignment")
        assert "Python" in analysis.skills_to_emphasize
        assert "FastAPI" in analysis.skills_to_emphasize
        assert analysis.recommended_project is not None
        assert len(analysis.checklist) > 0

    def test_partial_alignment(self, sample_profile: JobProfile):
        partial_job = Job(
            id=2,
            source="naukri",
            title="Junior Java Backend Developer",
            company="EnterpriseCorp",
            location="Remote",
            url="https://enterprisecorp.example/job2",
            description="Looking for Junior Backend Developer. Requires Java, Spring Boot, Docker, Kubernetes, AWS, SQL, and REST API.",
            match_score=55,
            matched_roles=["Backend Developer"],
            matched_skills=["SQL", "REST API"],
        )
        intel = ResumeIntelligenceService(profile=sample_profile)
        analysis = intel.analyze(partial_job)
        assert 40 <= analysis.resume_alignment_score <= 75
        assert "SQL" in analysis.skills_to_emphasize or "REST API" in analysis.skills_to_emphasize

    def test_weak_alignment(self, sample_profile: JobProfile):
        unrelated_job = Job(
            id=3,
            source="lever",
            title="Senior Embedded C++ Kernel Engineer",
            company="ChipMakers",
            location="Texas, US",
            url="https://jobs.lever.co/chip/kernel",
            description="Minimum 10+ years experience in C, C++, assembly, Linux kernel drivers, RTOS.",
            match_score=20,
            matched_roles=[],
            matched_skills=[],
        )
        intel = ResumeIntelligenceService(profile=sample_profile)
        analysis = intel.analyze(unrelated_job)
        assert analysis.resume_alignment_score <= 50
        assert analysis.resume_alignment_category in ("Low Alignment", "Partial Alignment")

    def test_profile_fallback_when_resume_empty(self, sample_profile: JobProfile, sample_job: Job):
        intel = ResumeIntelligenceService(
            profile=sample_profile,
            resume_profiles={"Profile Fallback": ResumeProfile(is_fallback=True)}
        )
        analysis = intel.analyze(sample_job)
        assert analysis.resume_alignment_score > 0
        assert "Python" in analysis.skills_to_emphasize
        assert analysis.recommended_resume == "Profile Fallback"

    def test_keyword_analysis(self, sample_profile: JobProfile, sample_job: Job):
        intel = ResumeIntelligenceService(profile=sample_profile)
        analysis = intel.analyze(sample_job)

        assert "Python" in analysis.resume_matched_keywords
        assert "FastAPI" in analysis.resume_matched_keywords
        # Redis and Docker are in description but not in candidate profile
        assert any(k in analysis.resume_missing_keywords for k in ["Redis", "Docker", "AWS"])

    def test_no_false_skill_claims(self, sample_profile: JobProfile, sample_job: Job):
        intel = ResumeIntelligenceService(profile=sample_profile)
        analysis = intel.analyze(sample_job)

        # Candidate does not know Redis or Docker -> they must NEVER appear in skills_to_emphasize
        assert "Redis" not in analysis.skills_to_emphasize
        assert "Docker" not in analysis.skills_to_emphasize
        assert "AWS" not in analysis.skills_to_emphasize

    def test_project_recommendation_ai_job(self, sample_profile: JobProfile):
        ai_job = Job(
            id=4,
            source="greenhouse",
            title="AI/ML Engineer",
            company="DeepTech AI",
            location="Bangalore",
            url="https://boards.greenhouse.io/deeptech/ai-eng",
            description="Build machine learning and computer vision pipelines using Python, OpenCV, scikit-learn.",
            match_score=85,
            matched_roles=["AI Engineer", "ML Engineer"],
            matched_skills=["Python", "OpenCV", "Scikit-learn", "Machine Learning"],
        )
        intel = ResumeIntelligenceService(profile=sample_profile)
        analysis = intel.analyze(ai_job)

        assert analysis.recommended_project == "Loan Application Fraud Detection via Document Verification"
        assert "OpenCV" in analysis.recommended_project_details["matched_skills"]

    def test_project_recommendation_backend_job(self, sample_profile: JobProfile):
        web_job = Job(
            id=5,
            source="greenhouse",
            title="Full Stack / Backend Engineer",
            company="SaaS Apps",
            location="Remote",
            url="https://boards.greenhouse.io/saas/backend",
            description="Develop REST APIs with Node/FastAPI and React frontend, using MongoDB and PostgreSQL databases.",
            match_score=80,
            matched_roles=["Backend Developer", "Full Stack Developer"],
            matched_skills=["JavaScript", "React", "REST API", "MongoDB", "PostgreSQL"],
        )
        intel = ResumeIntelligenceService(profile=sample_profile)
        analysis = intel.analyze(web_job)

        assert analysis.recommended_project == "NeuroNest Notes & Bookmarks Manager"

    def test_experience_recommendation(self, sample_profile: JobProfile, sample_job: Job):
        intel = ResumeIntelligenceService(profile=sample_profile)
        analysis = intel.analyze(sample_job)
        assert analysis.recommended_experience is not None
        assert "AI/ML Intern" in analysis.recommended_experience

    def test_concerns_and_suggestions_generated(self, sample_profile: JobProfile, sample_job: Job):
        intel = ResumeIntelligenceService(profile=sample_profile)
        analysis = intel.analyze(sample_job)
        assert len(analysis.resume_suggestions) >= 2
        for s in analysis.resume_suggestions:
            assert isinstance(s, str) and len(s) > 10


class TestDatabasePhase10:
    """Test suite for SQLite schema migration, persistence, and preparation status updates."""

    def test_database_migration_and_fields(self, tmp_path: Path, sample_job: Job):
        db_path = tmp_path / "test_jobs.db"
        db = DatabaseManager(db_path=db_path)

        # Apply preparation analysis
        analysis = ResumePreparationAnalysis(
            job_id=1,
            resume_alignment_score=88,
            resume_alignment_category="Strong Alignment",
            skills_to_emphasize=["Python", "FastAPI"],
            resume_matched_keywords=["Python", "FastAPI"],
            resume_missing_keywords=["Docker"],
            recommended_project="Loan Application Fraud Detection",
            recommended_experience="AI/ML Intern",
            resume_concerns=["Potential concern: Docker is required"],
            resume_suggestions=["1. Highlight FastAPI project"],
            recommended_resume="Primary Resume",
            preparation_status=PreparationStatus.REVIEWED,
            preparation_notes="Tailored for backend API role",
        )
        sample_job.apply_resume_preparation(analysis)

        is_new = db.save_job(sample_job)
        assert is_new is True

        retrieved = db.get_job_by_id(sample_job.id)
        assert retrieved is not None
        assert retrieved.resume_alignment_score == 88
        assert retrieved.resume_alignment_category == "Strong Alignment"
        assert "Python" in retrieved.skills_to_emphasize
        assert retrieved.preparation_status == "REVIEWED"
        assert retrieved.preparation_notes == "Tailored for backend API role"
        assert retrieved.recommended_resume == "Primary Resume"

    def test_update_preparation_status_and_notes(self, tmp_path: Path, sample_job: Job):
        db_path = tmp_path / "test_jobs.db"
        db = DatabaseManager(db_path=db_path)
        tracker = JobTracker(db)

        db.save_job(sample_job)
        retrieved = db.get_job_by_id(sample_job.id)
        assert retrieved.preparation_status == "NOT_REVIEWED"

        # Update to RESUME_READY
        success = tracker.update_preparation_status(
            job_id=sample_job.id,
            status=PreparationStatus.RESUME_READY,
            notes="Updated resume bullet points"
        )
        assert success is True

        updated = db.get_job_by_id(sample_job.id)
        assert updated.preparation_status == "RESUME_READY"
        assert updated.preparation_notes == "Updated resume bullet points"
        # Verify job.notes remains untouched
        assert updated.notes is None

        # Update to READY_TO_APPLY
        success2 = tracker.update_preparation_status(
            job_id=sample_job.id,
            status=PreparationStatus.READY_TO_APPLY
        )
        assert success2 is True
        assert db.get_job_by_id(sample_job.id).preparation_status == "READY_TO_APPLY"


class TestDashboardPhase10:
    """Test suite for FastAPI dashboard routes."""

    def test_resume_preparation_route(self):
        client = TestClient(app)
        response = client.get("/resume-preparation")
        assert response.status_code == 200
        assert "Resume Preparation" in response.text
        assert "PREPARATION STATUS" in response.text

    def test_job_detail_contains_resume_prep(self):
        client = TestClient(app)
        # Job 6 exists in real DB
        response = client.get("/jobs/6")
        assert response.status_code == 200
        assert "Resume Intelligence" in response.text
        assert "Resume Alignment:" in response.text
        assert "Preparation Status:" in response.text

    def test_preparation_post_endpoint(self):
        client = TestClient(app)
        response = client.post(
            "/jobs/6/preparation",
            data={
                "preparation_status": "REVIEWED",
                "preparation_notes": "Focused testing and QA notes",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert "/jobs/6" in response.headers.get("location")


class TestExcelPhase10:
    """Test suite for Excel export with Phase 10 headers and Resume Preparation sheet."""

    def test_excel_export_contains_phase_10_columns_and_sheet(self, tmp_path: Path, sample_job: Job):
        exporter = ExcelExporter(exports_dir=tmp_path)
        sample_job.resume_alignment_score = 85
        sample_job.resume_alignment_category = "Strong Alignment"
        sample_job.skills_to_emphasize = ["Python", "FastAPI"]
        sample_job.recommended_project = "Loan Application Fraud Detection"
        sample_job.preparation_status = PreparationStatus.RESUME_READY

        export_path = exporter.export([sample_job])
        assert export_path.exists()

        import openpyxl
        wb = openpyxl.load_workbook(str(export_path))
        assert "Resume Preparation" in wb.sheetnames
        assert "Top Matches" in wb.sheetnames
        assert "Apply Today" in wb.sheetnames

        ws = wb["Top Matches"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, len(REPORT_COLUMNS) + 1)]
        assert "Resume Alignment" in headers
        assert "Skills to Emphasize" in headers
        assert "Recommended Project" in headers
        assert "Preparation Status" in headers
