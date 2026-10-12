"""Comprehensive unit and integration tests for Phase 13 Job Search Intelligence & Source Optimization."""

from datetime import datetime
from pathlib import Path
import pytest
from starlette.testclient import TestClient

from app.database.database import DatabaseManager
from app.models.job import ApplicationPriority, Job, JobStatus
from app.models.profile import JobProfile
from app.models.search_intelligence import (
    DataConfidence,
    RecommendationPriority,
    SearchGapType,
    SourceActionRecommendation,
    StrategyStatus,
)
from app.services.excel_exporter import ExcelExporter
from app.services.matcher import is_india_eligible_remote
from app.services.search_intelligence import SearchIntelligenceService
from app.dashboard.app import app


@pytest.fixture
def temp_db(tmp_path):
    """Fixture providing an isolated SQLite database."""
    db_file = tmp_path / "test_search_intel.db"
    db = DatabaseManager(str(db_file))
    db.initialize()
    return db


@pytest.fixture
def sample_profile():
    """Fixture providing a mock target profile."""
    return JobProfile(
        target_roles=[
            "Python Developer",
            "Backend Developer",
            "Software Engineer",
            "Full Stack Developer",
            "Web Developer",
            "ML Engineer",
            "AI Engineer",
            "SDE Intern",
        ],
        preferred_locations=[
            "Bangalore",
            "Bengaluru",
            "Delhi",
            "Noida",
            "Ahmedabad",
            "Remote",
            "India",
        ],
        skills=["Python", "FastAPI", "React", "SQL", "Git"],
    )


def _create_mock_job(
    title: str = "Software Engineer",
    company: str = "TechCorp",
    location: str = "Bangalore, India",
    source: str = "greenhouse",
    match_score: int = 80,
    readiness_score: int = 75,
    resume_score: int = 70,
    priority: ApplicationPriority = ApplicationPriority.HIGH,
    status: JobStatus = JobStatus.NEW,
    applied_at: datetime = None,
    matched_roles: list = None,
    matched_locations: list = None,
) -> Job:
    """Helper creating standardized mock job instances."""
    return Job(
        title=title,
        company=company,
        location=location,
        source=source,
        url=f"https://example.com/job/{datetime.now().timestamp()}_{title.replace(' ', '_')}_{location.replace(' ', '_')}",
        description=f"Job listing for {title} in {location}",
        match_score=match_score,
        application_readiness_score=readiness_score,
        resume_alignment_score=resume_score,
        application_priority=priority,
        status=status,
        applied_at=applied_at,
        matched_roles=matched_roles or [title],
        matched_locations=matched_locations or ["Bangalore"],
    )


def test_source_intelligence_aggregation_and_relevance(temp_db, sample_profile):
    """Test source-level metric aggregations, relevance yield, and recommendation classifications."""
    # Insert 10 jobs for Greenhouse (7 relevant, 3 not relevant)
    for i in range(7):
        temp_db.save_job(_create_mock_job(
            title=f"Software Engineer {i}",
            source="greenhouse",
            location="Bangalore, India",
            match_score=85,
            priority=ApplicationPriority.HIGH,
        ))
    for i in range(3):
        temp_db.save_job(_create_mock_job(
            title=f"Senior Director {i}",
            source="greenhouse",
            location="Berlin, Germany",
            match_score=45,
            priority=ApplicationPriority.SKIP,
        ))

    service = SearchIntelligenceService(db=temp_db, profile=sample_profile)
    report = service.get_search_intelligence_report()

    gh = next((s for s in report.sources if s.source.lower() == "greenhouse"), None)
    assert gh is not None
    assert gh.jobs_found == 10
    assert gh.relevant_jobs == 7
    assert gh.relevance_rate == 70.0
    assert gh.confidence == DataConfidence.ACTIONABLE_DATA
    assert gh.recommendation == SourceActionRecommendation.INCREASE_SEARCH_DEPTH
    assert gh.recommendation_priority in (RecommendationPriority.HIGH, RecommendationPriority.CRITICAL)


def test_small_data_protection_and_insufficient_data(temp_db, sample_profile):
    """Ensure sources with < 5 jobs are flagged INSUFFICIENT_DATA and cannot be declared best source."""
    # Insert only 2 jobs for Lever (both 100% relevant)
    temp_db.save_job(_create_mock_job(title="Software Engineer", source="lever", match_score=90))
    temp_db.save_job(_create_mock_job(title="Backend Developer", source="lever", match_score=85))

    service = SearchIntelligenceService(db=temp_db, profile=sample_profile)
    report = service.get_search_intelligence_report()

    lev = next((s for s in report.sources if s.source.lower() == "lever"), None)
    assert lev is not None
    assert lev.jobs_found == 2
    assert lev.relevant_jobs == 2
    # Small data protection checks
    assert lev.confidence == DataConfidence.INSUFFICIENT_DATA
    assert lev.source_quality_score is None
    assert lev.quality_label == "Insufficient data"
    assert lev.recommendation == SourceActionRecommendation.INSUFFICIENT_DATA
    assert lev.recommendation_priority == RecommendationPriority.MONITOR
    assert "Fewer than 5 jobs" in lev.recommendation_reason


def test_source_health_integration_blocked_source(temp_db, sample_profile):
    """Ensure a blocked/rate-limited source is categorized as MONITOR rather than bad quality."""
    # 6 jobs with 0 relevant
    for i in range(6):
        temp_db.save_job(_create_mock_job(title=f"Job {i}", source="naukri", match_score=40))

    service = SearchIntelligenceService(db=temp_db, profile=sample_profile)
    # Simulate health data indicating blocked/anti-bot
    health_mock = [{
        "source": "Naukri",
        "enabled": True,
        "endpoint": "naukri.com",
        "status": "BLOCKED",
        "collected": 0,
        "error": "HTTP 403 Forbidden: anti-bot protection active",
    }]
    raw_jobs = temp_db.get_all_jobs_raw()
    perf = service.compute_source_performance(raw_jobs, health_data=health_mock)
    naukri = next(s for s in perf if s.source.lower() == "naukri")

    assert naukri.reliability == "Blocked"
    assert naukri.recommendation == SourceActionRecommendation.MONITOR
    assert naukri.recommendation_priority == RecommendationPriority.MONITOR
    assert "blocked or rate limited" in naukri.recommendation_reason.lower()


def test_foreign_remote_versus_india_eligible_remote(temp_db, sample_profile):
    """Ensure remote opportunities are cleanly distinguished into India-eligible vs foreign-only."""
    # India-eligible remote
    temp_db.save_job(_create_mock_job(
        title="Software Engineer India Remote",
        location="Bangalore, India - Remote",
        source="greenhouse",
        match_score=85,
    ))
    # Foreign-only remote (Canada / Poland / UK)
    temp_db.save_job(_create_mock_job(
        title="Software Engineer Canada",
        location="Remote, Canada; Remote, Poland; Remote, United Kingdom",
        source="greenhouse",
        match_score=65,
    ))

    # Verify helper function directly
    assert is_india_eligible_remote("Bangalore, India - Remote") is True
    assert is_india_eligible_remote("Remote, Canada; Remote, Poland; Remote, United Kingdom") is False

    service = SearchIntelligenceService(db=temp_db, profile=sample_profile)
    report = service.get_search_intelligence_report()

    # Verify locations breakdown
    india_remote = next((loc for loc in report.locations if "India Eligible" in loc.location_name), None)
    assert india_remote is not None
    assert india_remote.jobs_found >= 1

    foreign_remote = next((loc for loc in report.locations if "Foreign Only" in loc.location_name), None)
    assert foreign_remote is not None
    assert foreign_remote.foreign_remote >= 1
    assert foreign_remote.priority == RecommendationPriority.LOW


def test_role_intelligence_aggregation_and_ranking(temp_db, sample_profile):
    """Verify target role aggregation, score averages, and priority rankings."""
    # Software Engineer: 4 jobs, high match
    for i in range(4):
        temp_db.save_job(_create_mock_job(title=f"Software Engineer {i}", match_score=88, readiness_score=82))
    # Python Developer: 2 jobs, moderate match
    for i in range(2):
        temp_db.save_job(_create_mock_job(title=f"Python Developer {i}", match_score=72, readiness_score=68))

    service = SearchIntelligenceService(db=temp_db, profile=sample_profile)
    report = service.get_search_intelligence_report()

    se = next(r for r in report.roles if r.role_name == "Software Engineer")
    py = next(r for r in report.roles if r.role_name == "Python Developer")

    assert se.jobs_found == 4
    assert se.relevant_jobs == 4
    assert se.avg_match_score == 88.0
    assert se.priority == RecommendationPriority.HIGH

    assert py.jobs_found == 2
    assert py.relevant_jobs == 2
    assert py.avg_match_score == 72.0
    assert py.priority == RecommendationPriority.MEDIUM

    # Software Engineer should be ranked above Python Developer
    se_idx = report.roles.index(se)
    py_idx = report.roles.index(py)
    assert se_idx < py_idx


def test_role_location_matrix(temp_db, sample_profile):
    """Test Role x Location matrix generation and combination ranking."""
    # 3 jobs for Software Engineer in Bangalore
    for i in range(3):
        temp_db.save_job(_create_mock_job(
            title=f"Software Engineer {i}",
            location="Bangalore, India",
            match_score=85,
            priority=ApplicationPriority.HIGH,
        ))
    # 1 job for Python Developer in Delhi
    temp_db.save_job(_create_mock_job(
        title="Python Developer",
        location="Delhi, India",
        match_score=70,
        priority=ApplicationPriority.NORMAL,
    ))

    service = SearchIntelligenceService(db=temp_db, profile=sample_profile)
    report = service.get_search_intelligence_report()

    assert len(report.role_locations) >= 2
    top_combo = report.role_locations[0]
    assert top_combo.role_name == "Software Engineer"
    assert "Bangalore" in top_combo.location_name
    assert top_combo.relevant_jobs == 3
    assert top_combo.priority == RecommendationPriority.HIGH
    assert top_combo.status_label == "Strong Opportunity"


def test_source_role_and_source_location_matrices(temp_db, sample_profile):
    """Test Source x Role and Source x Location assessment classifications."""
    for i in range(3):
        temp_db.save_job(_create_mock_job(
            title="Software Engineer",
            source="greenhouse",
            location="Bangalore, India",
            match_score=85,
        ))

    service = SearchIntelligenceService(db=temp_db, profile=sample_profile)
    report = service.get_search_intelligence_report()

    sr = next((x for x in report.source_roles if x.source.lower() == "greenhouse" and x.role_name == "Software Engineer"), None)
    assert sr is not None
    assert sr.jobs_found == 3
    assert sr.relevant_jobs == 3
    assert sr.assessment == "strong"

    sl = next((x for x in report.source_locations if x.source.lower() == "greenhouse" and "Bangalore" in x.location_name), None)
    assert sl is not None
    assert sl.relevant_jobs == 3
    assert sl.assessment == "strong"


def test_search_coverage_gaps_detection(temp_db, sample_profile):
    """Ensure missing profile roles or locations are returned as SEARCH_GAP with careful wording."""
    # Only insert 1 job for Software Engineer in Bangalore
    temp_db.save_job(_create_mock_job(
        title="Software Engineer",
        location="Bangalore, India",
        match_score=80,
    ))

    service = SearchIntelligenceService(db=temp_db, profile=sample_profile)
    report = service.get_search_intelligence_report()

    # Configured roles like 'ML Engineer', 'Web Developer', 'SDE Intern' should be in search_gaps
    gap_roles = [g.item_name for g in report.search_gaps if g.gap_type == SearchGapType.ROLE_GAP]
    assert "ML Engineer" in gap_roles
    assert "Web Developer" in gap_roles

    # Check careful phrasing requirement
    ml_gap = next(g for g in report.search_gaps if g.item_name == "ML Engineer")
    assert "Insufficient historical coverage" in ml_gap.suggestion
    assert "no jobs exist" not in ml_gap.suggestion.lower()


def test_search_budget_allocation_sums_to_100(temp_db, sample_profile):
    """Ensure budget allocations always sum to exactly 100.0% under all sample sizes."""
    # Case 1: Insufficient data (<5 jobs)
    temp_db.save_job(_create_mock_job(source="greenhouse"))
    service = SearchIntelligenceService(db=temp_db, profile=sample_profile)
    report = service.get_search_intelligence_report()

    total_alloc = sum(b.allocated_percentage for b in report.budget_allocations)
    assert round(total_alloc, 1) == 100.0
    for b in report.budget_allocations:
        assert b.confidence == DataConfidence.INSUFFICIENT_DATA

    # Case 2: Actionable dataset (>10 jobs)
    for i in range(12):
        temp_db.save_job(_create_mock_job(
            title=f"Dev {i}",
            source="greenhouse",
            location="Bangalore",
            match_score=85,
        ))
    report2 = service.get_search_intelligence_report()
    total_alloc2 = sum(b.allocated_percentage for b in report2.budget_allocations)
    assert round(total_alloc2, 1) == 100.0


def test_recommendation_safety_constraints(temp_db, sample_profile):
    """Verify intelligence layer is strictly read-only and never modifies configuration or state."""
    initial_target_roles = list(sample_profile.target_roles)
    initial_preferred_locs = list(sample_profile.preferred_locations)

    job = _create_mock_job(title="Software Engineer", match_score=95, status=JobStatus.NEW)
    temp_db.save_job(job)

    service = SearchIntelligenceService(db=temp_db, profile=sample_profile)
    report = service.get_search_intelligence_report()

    # 1. Profile remains completely unmodified
    assert sample_profile.target_roles == initial_target_roles
    assert sample_profile.preferred_locations == initial_preferred_locs

    # 2. Database job status was NOT auto-applied
    db_job = temp_db.get_job_by_id(job.id)
    assert db_job.status == JobStatus.NEW
    assert db_job.applied_at is None


def test_dashboard_search_intelligence_route():
    """Verify HTTP GET /search-intelligence endpoint returns 200 with required UI elements."""
    client = TestClient(app)
    response = client.get("/search-intelligence")
    assert response.status_code == 200
    html = response.text
    assert "Search Intelligence" in html
    assert "Strategy Status" in html
    assert "Data Confidence" in html
    assert "Today&#39;s Recommended Search Priorities" in html or "Today's Recommended Search Priorities" in html
    assert "Recommended Search Effort Allocation" in html
    assert "Source Performance &amp; Depth Recommendations" in html or "Source Performance & Depth Recommendations" in html


def test_excel_export_includes_search_intelligence_sheet(tmp_path):
    """Verify Excel export workbook generates 'Search Intelligence' sheet without breaking existing sheets."""
    exporter = ExcelExporter(exports_dir=tmp_path)
    job = _create_mock_job(title="Backend Engineer", source="greenhouse", match_score=85)
    export_path = exporter.export([job], output_path=tmp_path / "test_report.xlsx")

    assert export_path.exists()
    import openpyxl
    wb = openpyxl.load_workbook(str(export_path))
    sheet_names = wb.sheetnames

    # Check all existing sheets are preserved
    expected_sheets = [
        "New Matches",
        "Top Matches",
        "All Relevant Jobs",
        "Apply Today",
        "Resume Preparation",
        "Applications",
        "Interview Preparation",
        "Analytics",
        "Search Intelligence",
    ]
    for s_name in expected_sheets:
        assert s_name in sheet_names, f"Expected sheet '{s_name}' missing from workbook"

    ws_intel = wb["Search Intelligence"]
    # Check headers
    cell_val = ws_intel.cell(row=1, column=1).value
    assert cell_val == "Strategy Element"
    assert ws_intel.max_row >= 8
