"""Comprehensive unit and integration tests for Phase 14 Advanced Personal Job-Agent Workflow."""

from datetime import date, datetime, timedelta
from pathlib import Path
import pytest
from starlette.testclient import TestClient

from app.database.database import DatabaseManager
from app.models.job import ApplicationPriority, Job, JobStatus
from app.models.job_agent import (
    AgentActionPriority,
    AgentActionStatus,
    AgentActionType,
    DailyJobPlan,
    JobAgentRecommendation,
    JobAgentReport,
)
from app.models.profile import JobProfile
from app.services.excel_exporter import ExcelExporter
from app.services.job_agent import JobAgentService
from app.dashboard.app import app


@pytest.fixture
def temp_db(tmp_path):
    """Fixture providing an isolated SQLite database."""
    db_file = tmp_path / "test_job_agent.db"
    db = DatabaseManager(str(db_file))
    db.initialize()
    return db


@pytest.fixture
def sample_profile():
    """Fixture providing a mock candidate profile."""
    return JobProfile(
        target_roles=[
            "Python Developer",
            "Backend Developer",
            "Software Engineer",
            "Full Stack Developer",
            "Web Developer",
            "ML Engineer",
            "AI Engineer",
            "SDE Intern",
        ],
        preferred_locations=[
            "Bangalore",
            "Bengaluru",
            "Delhi",
            "Noida",
            "Ahmedabad",
            "Remote",
            "India",
        ],
        skills=["Python", "FastAPI", "React", "SQL", "Git"],
    )


def _create_mock_job(
    title: str = "Software Engineer",
    company: str = "TechCorp",
    location: str = "Bangalore, India",
    source: str = "greenhouse",
    match_score: int = 80,
    readiness_score: int = 75,
    resume_score: int = 70,
    priority: ApplicationPriority = ApplicationPriority.HIGH,
    status: JobStatus = JobStatus.NEW,
    applied_at: datetime = None,
    follow_up_date: str = None,
    interview_date: str = None,
    interview_status: str = None,
    preparation_status: str = None,
) -> Job:
    """Helper creating standardized mock job instances."""
    return Job(
        title=title,
        company=company,
        location=location,
        source=source,
        url=f"https://example.com/job/{datetime.now().timestamp()}_{title.replace(' ', '_')}_{company.replace(' ', '_')}",
        description=f"Job listing for {title} requiring Python, FastAPI, and SQL.",
        match_score=match_score,
        application_readiness_score=readiness_score,
        resume_alignment_score=resume_score,
        application_priority=priority,
        status=status,
        applied_at=applied_at,
        follow_up_date=follow_up_date,
        interview_date=interview_date,
        interview_status=interview_status,
        preparation_status=preparation_status,
        matched_roles=[title],
        matched_locations=["Bangalore"],
    )


def test_agent_action_generation_prioritization_overdue_follow_up(temp_db, sample_profile):
    """Test that overdue follow-up receives top URGENT prioritization over upcoming follow-ups."""
    today_str = "2026-10-10"
    yesterday_str = "2026-10-09"
    future_str = "2026-10-15"

    # Job 1 has an overdue follow-up
    j1 = _create_mock_job(
        title="Backend Dev Overdue",
        company="CompanyOverdue",
        status=JobStatus.APPLIED,
        follow_up_date=yesterday_str,
    )
    # Job 2 has an upcoming follow-up
    j2 = _create_mock_job(
        title="Backend Dev Future",
        company="CompanyFuture",
        status=JobStatus.APPLIED,
        follow_up_date=future_str,
    )
    temp_db.save_job(j1)
    temp_db.save_job(j2)

    service = JobAgentService(db=temp_db, profile=sample_profile)
    plan = service.get_daily_plan(as_of_date=today_str)

    urgent_actions = [a for a in plan.urgent_actions if a.action_type == AgentActionType.FOLLOW_UP]
    assert len(urgent_actions) >= 1
    assert "CompanyOverdue" in urgent_actions[0].title
    assert urgent_actions[0].priority == AgentActionPriority.URGENT


def test_agent_interview_preparation_prioritization(temp_db, sample_profile):
    """Test that imminent interview is classified as URGENT while distant interview is HIGH."""
    today_str = "2026-10-10"
    imminent_date = "2026-10-11"  # Tomorrow (imminent)
    distant_date = "2026-10-25"   # 15 days out

    j_imminent = _create_mock_job(
        title="Python Engineer Imminent",
        company="ImminentCorp",
        status=JobStatus.INTERVIEW,
        interview_date=imminent_date,
        interview_status="SCHEDULED",
    )
    j_distant = _create_mock_job(
        title="Python Engineer Distant",
        company="DistantCorp",
        status=JobStatus.INTERVIEW,
        interview_date=distant_date,
        interview_status="SCHEDULED",
    )
    temp_db.save_job(j_imminent)
    temp_db.save_job(j_distant)

    service = JobAgentService(db=temp_db, profile=sample_profile)
    plan = service.get_daily_plan(as_of_date=today_str)

    imminent_act = next(a for a in plan.actions if a.company == "ImminentCorp")
    distant_act = next(a for a in plan.actions if a.company == "DistantCorp")

    assert imminent_act.priority == AgentActionPriority.URGENT
    assert distant_act.priority == AgentActionPriority.HIGH
    assert imminent_act.action_type == AgentActionType.PREPARE_INTERVIEW
    assert distant_act.action_type == AgentActionType.PREPARE_INTERVIEW
    # Verify interview package checklist and topics were integrated
    assert len(imminent_act.checklist) > 0


def test_agent_resume_prep_and_manual_apply_workflows(temp_db, sample_profile):
    """Test recommendation workflow: PREPARE_RESUME for unprepared vs APPLY_MANUALLY for ready."""
    today_str = "2026-10-10"

    # High match job needing resume prep
    j_unprepared = _create_mock_job(
        title="High Match Unprepared",
        company="UnpreparedCorp",
        match_score=88,
        readiness_score=85,
        priority=ApplicationPriority.HIGH,
        status=JobStatus.NEW,
        preparation_status="NOT_STARTED",
    )
    # High match job ready to apply
    j_ready = _create_mock_job(
        title="High Match Ready",
        company="ReadyCorp",
        match_score=85,
        readiness_score=82,
        priority=ApplicationPriority.HIGH,
        status=JobStatus.NEW,
        preparation_status="READY",
    )
    temp_db.save_job(j_unprepared)
    temp_db.save_job(j_ready)

    service = JobAgentService(db=temp_db, profile=sample_profile)
    plan = service.get_daily_plan(as_of_date=today_str)

    unprep_act = next(a for a in plan.actions if a.company == "UnpreparedCorp")
    ready_act = next(a for a in plan.actions if a.company == "ReadyCorp")

    assert unprep_act.action_type == AgentActionType.PREPARE_RESUME
    assert ready_act.action_type == AgentActionType.APPLY_MANUALLY
    assert "UnpreparedCorp" in unprep_act.title
    assert "ReadyCorp" in ready_act.title
    # Check that skills to emphasize were computed
    assert len(unprep_act.skills_to_emphasize) > 0


def test_agent_action_ranking_score_dimensions(temp_db, sample_profile):
    """Test deterministic action-ranking algorithm dimensions: overdue beats upcoming, high quality beats low."""
    today_str = "2026-10-10"

    j_overdue = _create_mock_job(
        title="Overdue Application",
        company="AlphaCorp",
        status=JobStatus.APPLIED,
        follow_up_date="2026-10-05",
    )
    j_high_match = _create_mock_job(
        title="High Match Opportunity",
        company="BetaCorp",
        match_score=92,
        readiness_score=88,
        status=JobStatus.NEW,
        preparation_status="READY",
    )
    j_low_match = _create_mock_job(
        title="Low Match Opportunity",
        company="GammaCorp",
        match_score=50,
        readiness_score=45,
        status=JobStatus.NEW,
    )
    temp_db.save_job(j_overdue)
    temp_db.save_job(j_high_match)
    temp_db.save_job(j_low_match)

    service = JobAgentService(db=temp_db, profile=sample_profile)
    plan = service.get_daily_plan(as_of_date=today_str)

    # Overdue action should rank first
    first_action = plan.actions[0]
    assert first_action.company == "AlphaCorp"
    assert first_action.priority == AgentActionPriority.URGENT

    # High match opportunity should rank before low match opportunity
    beta_idx = next(i for i, a in enumerate(plan.actions) if a.company == "BetaCorp")
    gamma_idx = next(i for i, a in enumerate(plan.actions) if a.company == "GammaCorp")
    assert beta_idx < gamma_idx


def test_agent_daily_target_integration_no_fake_progress(temp_db, sample_profile):
    """Test that daily target calculations reflect exact database numbers without fake progress."""
    today_str = date.today().isoformat()

    # 1 job applied today
    j_applied_today = _create_mock_job(
        title="Applied Today Job",
        company="AppliedCorp",
        status=JobStatus.APPLIED,
        applied_at=datetime.now(),
    )
    # 1 job applied 3 days ago
    j_applied_past = _create_mock_job(
        title="Past Applied Job",
        company="PastCorp",
        status=JobStatus.APPLIED,
        applied_at=datetime.now() - timedelta(days=3),
    )
    temp_db.save_job(j_applied_today)
    temp_db.save_job(j_applied_past)

    service = JobAgentService(db=temp_db, profile=sample_profile)
    plan = service.get_daily_plan()

    assert plan.target_daily == 5
    assert plan.target_applied_today == 1
    assert plan.target_remaining == 4
    assert plan.target_progress_pct == 20.0


def test_agent_application_neglect_detection(temp_db, sample_profile):
    """Test neglect detection: application with no follow-up scheduled after >5 days generates attention action."""
    today_str = "2026-10-10"
    applied_old = datetime(2026, 10, 2)  # 8 days ago, no follow_up_date

    j_neglected = _create_mock_job(
        title="Neglected Job",
        company="NeglectedCorp",
        status=JobStatus.APPLIED,
        applied_at=applied_old,
        follow_up_date=None,
    )
    temp_db.save_job(j_neglected)

    service = JobAgentService(db=temp_db, profile=sample_profile)
    plan = service.get_daily_plan(as_of_date=today_str)

    neglected_act = next((a for a in plan.actions if a.company == "NeglectedCorp"), None)
    assert neglected_act is not None
    assert neglected_act.action_type == AgentActionType.FOLLOW_UP
    assert "Needs Attention" in neglected_act.title or "attention" in neglected_act.reason.lower()


def test_agent_report_summary_and_action_cards(temp_db, sample_profile):
    """Test unified agent report synthesis including summary cards and safety notices."""
    j1 = _create_mock_job(
        title="Python Dev Ready",
        company="ReadyCorp",
        match_score=85,
        readiness_score=80,
        status=JobStatus.NEW,
        preparation_status="READY",
    )
    j2 = _create_mock_job(
        title="Software Engineer Prep",
        company="PrepCorp",
        match_score=78,
        readiness_score=72,
        status=JobStatus.NEW,
        preparation_status="NOT_STARTED",
    )
    temp_db.save_job(j1)
    temp_db.save_job(j2)

    service = JobAgentService(db=temp_db, profile=sample_profile)
    report = service.get_agent_report()

    assert isinstance(report, JobAgentReport)
    assert report.summary.relevant_jobs == 2
    assert len(report.today_applications) == 1
    assert report.today_applications[0].company == "ReadyCorp"
    assert len(report.resume_prep_jobs) == 1
    assert report.resume_prep_jobs[0].company == "PrepCorp"
    assert len(report.safety_notices) >= 4


def test_agent_safety_guarantees_no_auto_apply_no_status_mutation(temp_db, sample_profile):
    """Test safety guarantees: plan generation and action status marking never alter job database records."""
    j = _create_mock_job(
        title="Safety Check Job",
        company="SafetyCorp",
        status=JobStatus.NEW,
        preparation_status="NOT_STARTED",
    )
    saved_id = temp_db.save_job(j)

    service = JobAgentService(db=temp_db, profile=sample_profile)
    plan = service.get_daily_plan()

    # Verify job record in database remained completely unchanged
    db_job = temp_db.get_job_by_id(saved_id)
    assert db_job.status == JobStatus.NEW
    assert db_job.preparation_status == "NOT_STARTED"
    assert db_job.applied_at is None

    # Marking action complete in agent service must NOT mutate DB
    act = next(a for a in plan.actions if a.company == "SafetyCorp")
    service.mark_action_status(act.id, AgentActionStatus.COMPLETED)

    # Re-verify DB job remains unchanged
    db_job_after = temp_db.get_job_by_id(saved_id)
    assert db_job_after.status == JobStatus.NEW
    assert db_job_after.preparation_status == "NOT_STARTED"


def test_dashboard_job_agent_route_and_action_post(temp_db, monkeypatch):
    """Test that /job-agent returns HTTP 200 with proper HTML and POST /job-agent/action works."""
    # Monkeypatch get_db in routes to use isolated temp_db
    import app.dashboard.routes as routes_mod
    monkeypatch.setattr(routes_mod, "get_db", lambda: temp_db)

    # Add a job to test_db
    j = _create_mock_job(
        title="Frontend Dev",
        company="WebCorp",
        match_score=82,
        readiness_score=78,
        status=JobStatus.NEW,
    )
    temp_db.save_job(j)

    client = TestClient(app)
    resp = client.get("/job-agent")
    assert resp.status_code == 200
    assert "Personal Job Agent" in resp.text
    assert "Today's Action Plan" in resp.text
    assert "WebCorp" in resp.text

    # Test POST /job-agent/action
    resp_post = client.post(
        "/job-agent/action",
        data={"action_id": "act_resume_1", "action_status": "COMPLETED"},
        follow_redirects=False,
    )
    assert resp_post.status_code == 303


def test_excel_export_includes_job_agent_sheet(temp_db, tmp_path):
    """Test that Excel export creates Sheet 10 'Job Agent' while preserving all 9 prior sheets."""
    j1 = _create_mock_job(
        title="Lead Python Developer",
        company="PyCorp",
        match_score=85,
        readiness_score=80,
        status=JobStatus.NEW,
    )
    j2 = _create_mock_job(
        title="Full Stack Engineer",
        company="StackCorp",
        match_score=75,
        status=JobStatus.APPLIED,
        follow_up_date="2026-10-15",
    )
    temp_db.save_job(j1)
    temp_db.save_job(j2)
    jobs = temp_db.get_all_jobs()

    exporter = ExcelExporter(exports_dir=str(tmp_path))
    file_path = exporter.export(jobs)

    assert file_path.exists()

    import openpyxl
    wb = openpyxl.load_workbook(file_path, data_only=True)

    expected_sheets = [
        "New Matches",
        "Top Matches",
        "All Relevant Jobs",
        "Apply Today",
        "Resume Preparation",
        "Applications",
        "Interview Preparation",
        "Analytics",
        "Search Intelligence",
        "Job Agent",
    ]

    for sheet_name in expected_sheets:
        assert sheet_name in wb.sheetnames, f"Sheet '{sheet_name}' missing from exported workbook"

    ws_agent = wb["Job Agent"]
    headers = [cell.value for cell in ws_agent[1]]
    assert "Action Priority" in headers
    assert "Action Type" in headers
    assert "Job" in headers
    assert "Company" in headers
    assert "Match Score" in headers
    assert "Readiness" in headers
    assert "Resume Alignment" in headers
    assert "Recommended Action" in headers
    assert "Reason" in headers
    assert ws_agent.max_row >= 2
